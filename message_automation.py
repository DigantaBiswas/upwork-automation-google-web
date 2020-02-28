
import time
from selenium import webdriver
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.common.keys import Keys

import docx
import openpyxl
import os
import re



def read_text_message():
    text = docx.Document('message.docx')
    full_text = []
    for para in text.paragraphs:
        full_text.append(para.text)

    return full_text

def read_excel_numbers():
    wb = openpyxl.load_workbook('contacts.xlsx')
    ws = wb.active
    list_numbers = []

    for row in ws.iter_rows():
        number = row[8].value
        print(number)
        # print(number)
        number = re.findall(r'\d', str(number)) 
        number = list(map(int, number))
        number = ''.join([str(elem) for elem in number])  
        if number:
            list_numbers.append(number)
    return list_numbers


def main():
    try:
        text_message = read_text_message()[0]
    except Exception as e:
        print(e)

    recipients = read_excel_numbers()

    # Define the URL's we will open and a few other variables 
    main_url = 'https://messages.google.com/web/conversations' # URL A
    # tab_url = 'https://www.google.com' # URL B
    chromedriver = 'chromedriver'
    # Open main window with URL A
    browser= webdriver.Chrome(chromedriver)
    browser.get(main_url)
    time.sleep(20)

    div_id = 0
    for recipient in recipients:
        wait = WebDriverWait(browser, 50)

        start_chat_btn = browser.find_element_by_xpath('/html/body/mw-app/div/main/mw-main-container/div[1]/mw-main-nav/div/mw-fab-link/a')
        start_chat_btn.click()

        time.sleep(1)
        number_input_box = browser.find_element_by_xpath('//*[@id="mat-chip-list-{}"]/div/input'.format(div_id))
        number_input_box.send_keys(str(recipient))
        number_input_box.send_keys(Keys.ENTER)

        time.sleep(3)
        message_field = browser.find_element_by_xpath('/html/body/mw-app/div/main/mw-main-container/div[1]/mw-conversation-container/div/div/mws-message-compose/div/div[2]/div/mws-autosize-textarea/textarea')
        message_field.send_keys(str(text_message))
        message_field.send_keys(Keys.ENTER)
        time.sleep(2)
        
        div_id+=1


if __name__ == '__main__':
    main()